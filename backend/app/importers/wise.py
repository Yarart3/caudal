"""Importador de Wise (DISENO, sección 6.2).

Hoja `All transactions` con 32 columnas. Fechas de Excel e importes numéricos. Solo se usan las
columnas de la tabla de DISENO; el resto (titular, tarjeta, columnas del fondo...) se ignora.
Cada pago con tarjeta vende participaciones del fondo monetario EN LA MISMA FILA: eso no es un
gasto aparte, así que esas columnas no se leen y la fila es un gasto normal.
"""

from collections import Counter
from dataclasses import replace
from decimal import Decimal

from openpyxl.workbook.workbook import Workbook

from app.importers.base import (
    Importer,
    ImportFormatError,
    ParsedTransaction,
    ParseResult,
    Row,
    RowError,
    UnreadableValueError,
    check_balance_chain,
    closing,
    data_rows,
    has_headers,
    header_map,
    missing_headers,
    parse_date,
    require_headers,
    row_error,
    show,
    to_money,
)

SHEET = "All transactions"
ID = "Documento de identidad"
DATE = "Fecha"
AMOUNT = "Cantidad"
CURRENCY = "Divisa"
DESCRIPTION = "Descripción"
MERCHANT = "Comercio"
PAYER = "Nombre del pagador"
PAYEE = "Nombre del beneficiario"
DIRECTION = "Tipo de transacción"  # DEBIT / CREDIT
DETAILS_TYPE = "Tipo de detalles de transacción"  # CARD, DEPOSIT, TRANSFER...
REFERENCE = "Referencia del pago"  # solo para detectar transferencias desde imagin
BALANCE = "Saldo corriente"  # solo para comprobar el saldo
EXCHANGE_FROM = "Cambiar desde"
EXCHANGE_TO = "Cambiar a"

DETECT = [ID, DATE, AMOUNT, CURRENCY]
REQUIRED = [ID, DATE, AMOUNT, CURRENCY, DESCRIPTION]
SUPPORTED_CURRENCY = "EUR"  # Fase 1: solo EUR (multidivisa en la Fase 5)


class WiseImporter(Importer):
    name = "wise"
    label = "Wise"
    required = REQUIRED

    def detect(self, workbook: Workbook) -> bool:
        return SHEET in workbook.sheetnames and has_headers(workbook[SHEET], DETECT)

    def missing_columns(self, workbook: Workbook) -> list[str]:
        if SHEET not in workbook.sheetnames:
            return list(REQUIRED)
        return missing_headers(workbook[SHEET], REQUIRED)

    def parse(self, workbook: Workbook) -> ParseResult:
        if SHEET not in workbook.sheetnames:
            raise ImportFormatError(f"No se encuentra la hoja «{SHEET}»")
        sheet = workbook[SHEET]
        headers = header_map(sheet)
        require_headers(headers, REQUIRED)
        result = ParseResult()
        seen_ids: Counter[str] = Counter()

        for row in data_rows(sheet, headers):
            parsed = self._parse_row(row, result.errors)
            if parsed is None:
                continue
            # Si un ID se repitiera en el mismo archivo, cada aparición tiene su propia huella
            seen_ids[parsed.external_id or ""] += 1
            count = seen_ids[parsed.external_id or ""]
            if count > 1:
                parsed = replace(parsed, external_id=f"{parsed.external_id}#{count}")
            result.transactions.append(parsed)

        result.warnings.extend(check_balance_chain(result.transactions, self.label))
        result.closing_balance, result.closing_date = closing(result.transactions)
        return result

    def _parse_row(self, row: Row, errors: list[RowError]) -> ParsedTransaction | None:
        external_id = row.text(ID)
        if external_id is None:
            errors.append(RowError(row.number, ID, "", f"Fila {row.number}: falta el ID de Wise"))
            return None
        currency = row.text(CURRENCY)
        if currency is None or currency.upper() != SUPPORTED_CURRENCY:
            errors.append(
                RowError(
                    row.number,
                    CURRENCY,
                    currency or "",
                    f"Fila {row.number}: de momento solo se importan movimientos en EUR",
                )
            )
            return None
        try:
            when = parse_date(row.get(DATE))
        except UnreadableValueError as problem:
            errors.append(row_error(row.number, DATE, row.get(DATE), "la fecha", problem))
            return None
        try:
            amount = to_money(row.get(AMOUNT))
        except UnreadableValueError as problem:
            errors.append(row_error(row.number, AMOUNT, row.get(AMOUNT), "el importe", problem))
            return None

        direction = (row.text(DIRECTION) or "").upper()
        if (direction == "DEBIT" and amount > 0) or (direction == "CREDIT" and amount < 0):
            errors.append(
                RowError(
                    row.number,
                    AMOUNT,
                    show(row.get(AMOUNT)),
                    f"Fila {row.number}: el signo del importe no coincide con «{direction}»",
                )
            )
            return None

        balance: Decimal | None = None
        if row.get(BALANCE) is not None:
            try:
                balance = to_money(row.get(BALANCE))
            except UnreadableValueError:
                balance = None  # solo sirve para comprobar el saldo: no impide importar la fila

        description = _description(row, amount)
        if description is None:
            errors.append(
                RowError(row.number, DESCRIPTION, "", f"Fila {row.number}: el concepto está vacío")
            )
            return None

        return ParsedTransaction(
            row=row.number,
            date=when,
            description=description,
            amount=amount,
            balance=balance,
            external_id=external_id,
            transfer_hint=_transfer_hint(row, direction),
        )


def _description(row: Row, amount: Decimal) -> str | None:
    """Comercio → beneficiario (gasto) o pagador (ingreso) → Descripción."""
    name = row.text(PAYEE) if amount < 0 else row.text(PAYER)
    return row.text(MERCHANT) or name or row.text(DESCRIPTION)


def _transfer_hint(row: Row, direction: str) -> str | None:
    if row.get(EXCHANGE_FROM) is not None or row.get(EXCHANGE_TO) is not None:
        return "conversion"
    reference = (row.text(REFERENCE) or "").casefold()
    details = (row.text(DETAILS_TYPE) or "").upper()
    if direction == "CREDIT" and details == "DEPOSIT" and "imaginbank" in reference:
        return "imagin"
    return None
