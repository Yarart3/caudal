"""Piezas comunes de los importadores de extractos.

Un importador recibe el Excel ya abierto, comprueba si es de su banco (`detect`) y lo
convierte en movimientos (`parse`). No toca la BD: guardar es cosa del servicio de
importación (Fase 1b, parte 1.2). Así se puede probar con los ejemplos de `samples/` sin nada más.
"""

import re
import unicodedata
from abc import ABC, abstractmethod
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO
from typing import Any
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

CENT = Decimal("0.01")


class ImportFormatError(Exception):
    """El archivo no se puede importar entero (no es un Excel, falta una columna...).

    El mensaje está en castellano y se enseña tal cual al usuario.
    """


@dataclass(frozen=True)
class RowError:
    """Una fila que no se ha podido leer. No aborta el resto de la importación."""

    row: int  # número de fila del Excel (la 1 son las cabeceras)
    field: str  # columna que ha fallado, con el nombre del extracto ("Importe")
    value: str  # el dato tal cual venía
    message: str  # explicación en castellano


@dataclass(frozen=True)
class ParsedTransaction:
    """Un movimiento leído del extracto, todavía sin guardar."""

    row: int
    date: date
    description: str
    amount: Decimal  # negativo = gasto, positivo = ingreso
    balance: Decimal | None = None  # saldo tras el movimiento, según el extracto (no se guarda)
    external_id: str | None = None  # ID del banco, si lo da (Wise: "Documento de identidad")
    # Pista de transferencia interna que da el propio extracto:
    #   "imagin"     → ingreso en Wise que viene de imagin ("imaginBank" en la referencia)
    #   "conversion" → conversión de divisa dentro de Wise
    transfer_hint: str | None = None


@dataclass
class ParseResult:
    transactions: list[ParsedTransaction] = field(default_factory=list)
    errors: list[RowError] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    # Saldo más reciente del extracto y su fecha (para el saldo de referencia de la cuenta)
    closing_balance: Decimal | None = None
    closing_date: date | None = None


class Importer(ABC):
    """Interfaz común. Cada banco implementa `detect` y `parse`."""

    name: str  # identificador que se guarda en accounts.import_format ("imagin", "wise")
    label: str  # nombre para los mensajes ("imagin", "Wise")
    required: list[str]  # columnas sin las que no se puede importar

    @abstractmethod
    def detect(self, workbook: Workbook) -> bool:
        """True si el archivo tiene la forma del extracto de este banco."""

    @abstractmethod
    def missing_columns(self, workbook: Workbook) -> list[str]:
        """Columnas necesarias que no están (para explicar por qué no se reconoce el archivo)."""

    @abstractmethod
    def parse(self, workbook: Workbook) -> ParseResult:
        """Lee los movimientos. Lanza ImportFormatError si falta una columna necesaria."""


# ---------- Abrir el archivo ----------


def open_workbook(data: bytes) -> Workbook:
    """Abre un .xlsx desde memoria. `data_only` da el valor calculado de las fórmulas."""
    try:
        return load_workbook(BytesIO(data), read_only=True, data_only=True)
    except (BadZipFile, InvalidFileException, KeyError, OSError, ValueError):
        raise ImportFormatError("El archivo no es un Excel (.xlsx) válido") from None


# ---------- Cabeceras y filas ----------


def _header_key(value: Any) -> str:
    return " ".join(str(value).split()).casefold() if value is not None else ""


def header_map(sheet: Worksheet) -> dict[str, int]:
    """{nombre de columna normalizado: índice} de la fila 1. Las columnas se leen por nombre."""
    first = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
    return {_header_key(v): i for i, v in enumerate(first) if _header_key(v)}


def has_headers(sheet: Worksheet, names: list[str]) -> bool:
    headers = header_map(sheet)
    return all(_header_key(n) in headers for n in names)


def missing_headers(sheet: Worksheet, names: list[str]) -> list[str]:
    headers = header_map(sheet)
    return [n for n in names if _header_key(n) not in headers]


def require_headers(headers: dict[str, int], names: list[str]) -> None:
    for name in names:
        if _header_key(name) not in headers:
            raise ImportFormatError(f"No se encuentra la columna «{name}»")


class Row:
    """Acceso a una fila por nombre de columna. Celda vacía y texto vacío ('') son lo mismo."""

    def __init__(self, number: int, values: tuple[Any, ...], headers: dict[str, int]):
        self.number = number
        self._values = values
        self._headers = headers

    def get(self, name: str) -> Any:
        index = self._headers.get(_header_key(name))
        if index is None or index >= len(self._values):
            return None
        value = self._values[index]
        if isinstance(value, str) and not value.strip():
            return None
        return value

    def text(self, name: str) -> str | None:
        value = self.get(name)
        return None if value is None else " ".join(str(value).split())

    def is_blank(self) -> bool:
        return all(v is None or (isinstance(v, str) and not v.strip()) for v in self._values)


def data_rows(sheet: Worksheet, headers: dict[str, int]) -> list[Row]:
    """Filas de datos (de la 2 en adelante), sin las que están completamente vacías."""
    rows = []
    for number, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        row = Row(number, values, headers)
        if not row.is_blank():
            rows.append(row)
    return rows


# ---------- Normalización de valores ----------


class UnreadableValueError(ValueError):
    """Un dato concreto no se puede leer. El mensaje va al usuario."""


_DATE_TEXT = re.compile(r"^(\d{1,2})/(\d{1,2})/(\d{4})$")


def parse_date(value: Any) -> date:
    """Fecha como texto `dd/mm/aaaa` o como fecha de Excel (las dos formas aparecen)."""
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, int | float) and not isinstance(value, bool):
        # Número de serie de Excel sin formato de fecha (por ejemplo, 46294)
        if 1 <= value < 2958466:  # 1900-01-01 .. 9999-12-31
            converted = from_excel(value)
            if isinstance(converted, datetime):
                return converted.date()
        raise UnreadableValueError("no es una fecha válida")
    if isinstance(value, str):
        match = _DATE_TEXT.match(value.strip())
        if match:
            day, month, year = (int(g) for g in match.groups())
            try:
                return date(year, month, day)
            except ValueError:
                pass
    raise UnreadableValueError("no es una fecha con el formato dd/mm/aaaa")


# Importe español: miles con punto (opcionales), decimales con coma (0, 1 o 2 cifras)
_ES_AMOUNT = re.compile(r"^([+-]?)(\d{1,3}(?:\.\d{3})+|\d+)(?:,(\d{1,2}))?$")


def parse_es_amount(value: Any, currency: str = "EUR") -> Decimal:
    """Importe como lo da imagin: `-14,50EUR`, `-23EUR`, `-0,5EUR`, `1.234,56EUR` → Decimal.

    Si algún día llega como número de Excel, también se acepta.
    """
    if isinstance(value, int | float | Decimal) and not isinstance(value, bool):
        return to_money(value)
    if not isinstance(value, str):
        raise UnreadableValueError("no es un importe")
    text = value.strip().replace("\u2212", "-")  # signo menos tipográfico
    text = "".join(text.split())  # quita espacios, también los de "-23 EUR"
    upper = text.upper()
    if upper.endswith(currency):
        text = text[: -len(currency)]
    elif upper.endswith("€"):
        text = text[:-1]
    elif re.search(r"[A-Z]{3}$", upper):
        raise UnreadableValueError(f"la divisa no es {currency}")
    match = _ES_AMOUNT.match(text)
    if not match:
        raise UnreadableValueError("no es un importe con el formato -1.234,56EUR")
    sign, integer, decimals = match.groups()
    number = Decimal(f"{sign}{integer.replace('.', '')}.{decimals or '0'}")
    return number.quantize(CENT)


def to_money(value: Any) -> Decimal:
    """Número de Excel (int o float) → Decimal con 2 decimales, sin pasar por errores de float.

    `str(float)` da la representación más corta (664.35 → "664.35"), así que la conversión es
    exacta para importes con 2 decimales. Si trae más decimales, no es un importe en euros.
    """
    if isinstance(value, bool):
        raise UnreadableValueError("no es un importe")
    if isinstance(value, str):
        value = value.strip().replace("\u2212", "-")
    try:
        number = Decimal(str(value))
    except InvalidOperation:
        raise UnreadableValueError("no es un importe") from None
    if not number.is_finite():
        raise UnreadableValueError("no es un importe")
    quantized = number.quantize(CENT)
    if quantized != number:
        raise UnreadableValueError("tiene más de dos decimales")
    if abs(quantized) >= Decimal("1e10"):  # NUMERIC(12,2)
        raise UnreadableValueError("es demasiado grande")
    return quantized


def normalize_text(text: str) -> str:
    """Mayúsculas, sin acentos y con un solo espacio. Para huellas y (en la 1c) reglas."""
    decomposed = unicodedata.normalize("NFKD", text)
    without_accents = "".join(c for c in decomposed if not unicodedata.combining(c))
    return " ".join(without_accents.upper().split())


def show(value: Any) -> str:
    """Dato tal cual para el mensaje de error."""
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%d/%m/%Y")
    return str(value)


def row_error(
    row: int, field: str, value: Any, what: str, problem: UnreadableValueError
) -> RowError:
    """Error de una fila con el formato de DISENO: «Fila 5: no se ha podido leer el importe…»."""
    shown = show(value)
    detail = f"«{shown}» ({problem})" if shown else "(está vacío)"
    return RowError(row, field, shown, f"Fila {row}: no se ha podido leer {what} {detail}")


# ---------- Comprobación del saldo ----------

_Gap = tuple[int, int, Decimal, Decimal]  # fila anterior, fila siguiente, esperado, extracto


def check_balance_chain(transactions: list[ParsedTransaction], label: str) -> list[str]:
    """Comprueba fila a fila que saldo anterior + importe = saldo del extracto.

    Los bancos dan el extracto del más reciente al más antiguo, pero no hay garantía: se prueba
    en los dos sentidos y se usa el que cuadra mejor. Cada salto es un aviso: puede faltar un
    movimiento (o el extracto tiene filas reordenadas).
    """
    # (fila, importe, saldo) de las filas que traen saldo, en el orden del archivo
    rows = [(t.row, t.amount, t.balance) for t in transactions if t.balance is not None]
    if len(rows) < 2:
        return []

    def gaps(ordered: list[tuple[int, Decimal, Decimal]]) -> list[_Gap]:
        found = []
        for (prev_row, _, prev_balance), (row, amount, balance) in zip(
            ordered, ordered[1:], strict=False
        ):
            if prev_balance + amount != balance:
                found.append((prev_row, row, prev_balance + amount, balance))
        return found

    newest_first = gaps(list(reversed(rows)))  # el archivo va del más reciente al más antiguo
    oldest_first = gaps(rows)
    found = newest_first if len(newest_first) <= len(oldest_first) else oldest_first
    return [
        f"El saldo de {label} no cuadra entre la fila {prev_row} y la fila {row}: "
        f"tendría que ser {format_eur(expected)} y el extracto dice {format_eur(balance)}. "
        "Puede faltar algún movimiento."
        for prev_row, row, expected, balance in found
    ]


def closing(transactions: list[ParsedTransaction]) -> tuple[Decimal | None, date | None]:
    """Saldo más reciente del extracto: el de la fecha más alta y, en empate, la fila de arriba
    (los extractos van del más reciente al más antiguo)."""
    with_balance = [t for t in transactions if t.balance is not None]
    if not with_balance:
        return None, None
    latest = min(with_balance, key=lambda t: (-t.date.toordinal(), t.row))
    return latest.balance, latest.date


def format_eur(amount: Decimal) -> str:
    """1234.5 → '1.234,50 €' (solo para mensajes)."""
    sign = "-" if amount < 0 else ""
    integer, decimals = f"{abs(amount):.2f}".split(".")
    groups = []
    while len(integer) > 3:
        groups.insert(0, integer[-3:])
        integer = integer[:-3]
    groups.insert(0, integer)
    return f"{sign}{'.'.join(groups)},{decimals} €"
